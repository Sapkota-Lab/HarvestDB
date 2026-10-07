from dataclasses import replace
from datetime import date
import os
from pathlib import Path
from shutil import copyfile

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.dialects.postgresql import JSONB
from sqlalchemy.ext.compiler import compiles
from sqlalchemy.orm import Session

from app.db.base import Base
from app.models.field import Field
from app.models.harvest_event import HarvestEvent
from app.models.harvest_record import HarvestRecord
from app.services.workbook_import import WorkbookImport, import_workbook, read_workbook

HEADERS = [
    "SN", "Genotypes", "Location", "Date of harvest", "Harvest",
    "Grade 1 Marketable (weight)", "Grade 1 Marketable (counts)", "Replicate",
]
ROW = [1, "UKYT-25-001", "High tunnel", date(2026, 7, 13), "First", 0, None, "N/A"]


@compiles(JSONB, "sqlite")
def sqlite_jsonb(_type, _compiler, **_kwargs):
    return "JSON"


@pytest.fixture
def db():
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    with Session(engine) as session:
        yield session
    engine.dispose()


def make_workbook(tmp_path, sheets=None, headers=None):
    path = tmp_path / "cleaned.xlsx"
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in (sheets or {"HT_01_First_2026-07-13": [ROW]}).items():
        sheet = workbook.create_sheet(name)
        sheet.append(headers or HEADERS)
        for row in rows:
            sheet.append(row)
    workbook.save(path)
    workbook.close()
    return path


def count_rows(db, model):
    return db.scalar(select(func.count()).select_from(model))


def test_preserves_source_values_and_uses_sn_as_temporary_plot(tmp_path):
    plan = read_workbook(make_workbook(tmp_path))
    row = plan.rows[0]
    assert row.plot_number == "1"
    assert row.location == "High Tunnel"
    assert row.harvest_date == date(2026, 7, 13)
    assert row.harvest_number == 1
    assert row.dynamic_data["serial_number"] == 1
    assert row.dynamic_data["genotype"] == "UKYT-25-001"
    assert row.dynamic_data["grade_1_marketable_weight"] == 0
    assert row.dynamic_data["grade_1_marketable_count"] is None
    assert row.dynamic_data["replicate"] == "N/A"
    assert row.source_key == ("HT_01_First_2026-07-13", 2)


def test_preserves_inferred_date_reasoning_and_uncached_formulas(tmp_path):
    path = make_workbook(tmp_path, headers=HEADERS + ["Total Weight (g)"],
                         sheets={"HT_01_First_2026-07-13": [ROW + ["=F2"]]})
    workbook = load_workbook(path)
    index = workbook.create_sheet("_Sheet_Index")
    index.append(["New sheet name", "Original name", "Location", "Harvest", "Date assigned", "Date reasoning"])
    index.append(["HT_01_First_2026-07-13", "First", "High tunnel", "First", "2026-07-13", "Inferred from weekly cadence"])
    workbook.save(path)
    workbook.close()
    plan = read_workbook(path)
    assert plan.uncached_formulas == 1
    assert plan.rows[0].dynamic_data["total_weight_g"] is None
    source = plan.rows[0].dynamic_data["_source"]
    assert source["formulas"] == {"Total Weight (g)": "=F2"}
    assert source["date_reasoning"] == "Inferred from weekly cadence"


def test_excludes_undated_and_non_harvest_sheets(tmp_path):
    path = make_workbook(tmp_path, sheets={
        "HT_01_First_2026-07-13": [ROW],
        "SF_Hybrid": [[1, "G1", "South Farm", "N/A", "N/A", 10, 2, "R1"]],
        "Lab_Specific_Gravity": [ROW],
        "HT_Entry_Metadata": [ROW],
    })
    plan = read_workbook(path)
    assert len(plan.rows) == 1
    assert len(plan.excluded_sheets) == 3
    assert plan.excluded_sheets["SF_Hybrid"] == "No harvest date recorded"


@pytest.mark.parametrize("column,value", [
    (0, None), (0, 1.5), (0, True), (0, "a" * 65), (0, "N/A"),
    (2, "Unknown site"), (3, "N/A"), (3, date(2027, 7, 13)), (4, "Unknown harvest"),
])
def test_rejects_invalid_event_or_plot_data(tmp_path, column, value):
    row = list(ROW)
    row[column] = value
    with pytest.raises(ValueError, match="row 2"):
        read_workbook(make_workbook(tmp_path, sheets={"HT_01_First_2026-07-13": [row]}))


def test_unknown_columns_are_not_silently_dropped(tmp_path):
    path = make_workbook(tmp_path, headers=HEADERS + ["New measurement"],
                         sheets={"HT_01_First_2026-07-13": [ROW + [12]]})
    with pytest.raises(ValueError, match="unknown columns"):
        read_workbook(path)


def test_unknown_genotype_is_preserved_without_inventing_an_identifier(tmp_path):
    row = list(ROW)
    row[1] = "N/A"
    plan = read_workbook(make_workbook(tmp_path, sheets={"HT_01_First_2026-07-13": [row]}))
    assert plan.summary()["records_without_genotype"] == 1
    assert plan.rows[0].dynamic_data["genotype"] == "N/A"


def test_same_sn_in_separate_sheets_is_not_deduplicated(tmp_path, db):
    row = [1, "G1", "South Farm", date(2026, 8, 18), "Third", 12, 3, "R1"]
    other = list(row)
    other[1] = "G2"
    plan = read_workbook(make_workbook(tmp_path, sheets={
        "SF_03_Third_2026-08-18_Nam": [row],
        "SF_03_Third_2026-08-18_Ella": [other],
    }))
    with db.begin():
        result = import_workbook(db, plan)
    assert result["records_created"] == 2
    assert count_rows(db, HarvestEvent) == 1
    assert {record.dynamic_data["genotype"] for record in db.scalars(select(HarvestRecord))} == {"G1", "G2"}


def test_row_dates_take_priority_over_the_sheet_name(tmp_path, db):
    second_day = list(ROW)
    second_day[0] = 2
    second_day[3] = date(2026, 7, 14)
    plan = read_workbook(make_workbook(tmp_path, sheets={"HT_01_First_2026-07-13": [ROW, second_day]}))
    with db.begin():
        result = import_workbook(db, plan)
    assert result["events_created"] == 2
    assert {event.harvest_date for event in db.scalars(select(HarvestEvent))} == {date(2026, 7, 13), date(2026, 7, 14)}


def test_repeat_import_and_renamed_file_create_no_duplicates(tmp_path, db):
    path = make_workbook(tmp_path)
    plan = read_workbook(path)
    with db.begin():
        first = import_workbook(db, plan)
    renamed = tmp_path / "renamed.xlsx"
    copyfile(path, renamed)
    with db.begin():
        second = import_workbook(db, read_workbook(renamed))
    assert first == {"fields_created": 1, "events_created": 1, "records_created": 1, "records_skipped": 0}
    assert second == {"fields_created": 0, "events_created": 0, "records_created": 0, "records_skipped": 1}
    assert count_rows(db, HarvestRecord) == 1


def test_changed_source_row_stops_import_without_overwriting_data(tmp_path, db):
    plan = read_workbook(make_workbook(tmp_path))
    with db.begin():
        import_workbook(db, plan)
    changed_data = {**plan.rows[0].dynamic_data, "genotype": "Different genotype"}
    changed = replace(plan.rows[0], dynamic_data=changed_data)
    with pytest.raises(ValueError, match="Previously imported row changed"), db.begin():
        import_workbook(db, WorkbookImport(rows=[changed]))
    assert db.scalar(select(HarvestRecord)).dynamic_data["genotype"] == "UKYT-25-001"


def test_existing_fields_and_events_are_reused(tmp_path, db):
    with db.begin():
        location = Field(name="High tunnel")
        db.add(location)
        db.flush()
        db.add(HarvestEvent(field_id=location.id, harvest_date=date(2026, 7, 13), harvest_number=1))
    with db.begin():
        result = import_workbook(db, read_workbook(make_workbook(tmp_path)))
    assert result["fields_created"] == result["events_created"] == 0
    assert count_rows(db, Field) == count_rows(db, HarvestEvent) == 1


def test_failure_rolls_back_all_new_fields_events_and_records(tmp_path, db):
    with db.begin():
        location = Field(name="South Farm")
        db.add(location)
        db.flush()
        for _ in range(2):
            db.add(HarvestEvent(field_id=location.id, harvest_date=date(2026, 8, 18), harvest_number=3))
    plan = read_workbook(make_workbook(tmp_path, sheets={
        "HT_01_First_2026-07-13": [ROW],
        "SF_03_Third_2026-08-18_Nam": [[1, "G1", "South Farm", date(2026, 8, 18), "Third", 12, 3, "R1"]],
    }))
    with pytest.raises(ValueError, match="Ambiguous harvest event"), db.begin():
        import_workbook(db, plan)
    assert count_rows(db, HarvestRecord) == 0
    assert count_rows(db, Field) == 1
    assert count_rows(db, HarvestEvent) == 2


def test_actual_workbook_when_available(db):
    source_path = os.environ.get("HARVEST_TEST_WORKBOOK")
    if not source_path:
        pytest.skip("Set HARVEST_TEST_WORKBOOK to test the full source workbook")
    path = Path(source_path)
    plan = read_workbook(path)
    assert len(plan.rows) == 601
    assert plan.summary()["fields"] == {"High Tunnel": 214, "South Farm": 304, "Robinson": 83}
    assert plan.summary()["harvest_events"] == 14
    with db.begin():
        result = import_workbook(db, plan)
    assert result["records_created"] == 601
    with db.begin():
        repeated = import_workbook(db, plan)
    assert repeated["records_created"] == 0
    assert repeated["records_skipped"] == 601
    assert count_rows(db, HarvestRecord) == 601
    source_rows = {row.source_key: row for row in plan.rows}
    for record, event in db.execute(
        select(HarvestRecord, HarvestEvent).join(HarvestEvent, HarvestRecord.harvest_event_id == HarvestEvent.id)
    ):
        source = record.dynamic_data["_source"]
        expected = source_rows[(source["sheet"], source["row"])]
        assert record.dynamic_data == expected.dynamic_data
        assert record.plot_number == expected.plot_number
        assert event.harvest_date == expected.harvest_date
        assert event.harvest_number == expected.harvest_number
