import math
import re
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select, text
from sqlalchemy.orm import Session

from app.models.field import Field
from app.models.harvest_event import HarvestEvent
from app.models.harvest_record import HarvestRecord

DATASET = "cleaned-tomato-harvest-2026"
LOCATIONS = {"high tunnel": "High Tunnel", "south farm": "South Farm", "robinson": "Robinson"}
HARVEST_NUMBERS = {
    "first": 1, "second": 2, "third": 3, "fourth": 4,
    "fifth": 5, "sixth": 6, "seventh": 7,
}
COLUMNS = {
    "SN": "serial_number",
    "Genotypes": "genotype",
    "Location": "location",
    "Harvest": "harvest",
    "Experiment": "experiment",
    "Grade 1 Marketable (weight)": "grade_1_marketable_weight",
    "Grade 1 Marketable (counts)": "grade_1_marketable_count",
    "Grade 2 Marketable (weight)": "grade_2_marketable_weight",
    "Grade 2 Marketable (counts)": "grade_2_marketable_count",
    "Unmarketables (weight)": "unmarketable_weight",
    "Unmarketables (counts)": "unmarketable_count",
    "Weight of 20 marketable (g)": "weight_of_20_marketable_g",
    "Weight of 20 fruit (g)": "weight_of_20_fruit_g",
    "Total Weight (g)": "total_weight_g",
    "Total Weight (kg)": "total_weight_kg",
    "Fruit count": "fruit_count",
    "Calculated": "calculated",
    "Note": "note",
    "Genotype_raw": "genotype_raw",
    "Replicate": "replicate",
    "Entered_by": "entered_by",
}
REQUIRED_COLUMNS = {"SN", "Genotypes", "Location", "Date of harvest", "Harvest"}


@dataclass
class WorkbookRow:
    location: str
    harvest_date: date
    harvest_number: int
    plot_number: str
    dynamic_data: dict

    @property
    def source_key(self) -> tuple[str, int]:
        source = self.dynamic_data["_source"]
        return source["sheet"], source["row"]

    @property
    def event_key(self) -> tuple[str, date, int]:
        return self.location, self.harvest_date, self.harvest_number


@dataclass
class WorkbookImport:
    rows: list[WorkbookRow] = field(default_factory=list)
    excluded_sheets: dict[str, str] = field(default_factory=dict)
    uncached_formulas: int = 0

    def summary(self) -> dict:
        return {
            "records": len(self.rows),
            "fields": dict(Counter(row.location for row in self.rows)),
            "harvest_events": len({row.event_key for row in self.rows}),
            "uncached_formulas": self.uncached_formulas,
            "records_without_genotype": sum(
                row.dynamic_data.get("genotype") in (None, "", "N/A") for row in self.rows
            ),
            "excluded_sheets": self.excluded_sheets,
        }


def read_workbook(path: Path) -> WorkbookImport:
    plan = WorkbookImport()
    workbook = load_workbook(path, read_only=True, data_only=False)
    cached = load_workbook(path, read_only=True, data_only=True)
    try:
        index = {}
        if "_Sheet_Index" in workbook.sheetnames:
            for row_number, values in enumerate(
                workbook["_Sheet_Index"].iter_rows(min_row=2, values_only=True), start=2
            ):
                label = f"_Sheet_Index row {row_number}"
                index[values[0]] = {
                    "original_sheet": _json_value(values[1], label),
                    "date_assigned": _json_value(values[4], label),
                    "date_reasoning": _json_value(values[5], label),
                }

        for sheet in workbook:
            if not re.match(r"^(HT|SF|RC)_\d{2}_", sheet.title):
                plan.excluded_sheets[sheet.title] = (
                    "No harvest date recorded" if sheet.title == "SF_Hybrid"
                    else "Metadata, cleaning history, or lab assay sheet"
                )
                continue

            source_rows = sheet.iter_rows()
            cached_rows = cached[sheet.title].iter_rows()
            headers = [cell.value for cell in next(source_rows)]
            next(cached_rows)
            if len(headers) != len(set(headers)) or any(not header for header in headers):
                raise ValueError(f"{sheet.title}: headers must be populated and unique")
            missing = REQUIRED_COLUMNS - set(headers)
            unknown = set(headers) - set(COLUMNS) - {"Date of harvest"}
            if missing or unknown:
                raise ValueError(f"{sheet.title}: missing columns {sorted(missing)}, unknown columns {sorted(unknown)}")

            for row_number, (source_cells, cached_cells) in enumerate(zip(source_rows, cached_rows), start=2):
                if all(cell.value is None for cell in source_cells):
                    continue
                label = f"{sheet.title} row {row_number}"
                values = {}
                formulas = {}
                for header, cell, cached_cell in zip(headers, source_cells, cached_cells):
                    value = cell.value
                    if cell.data_type == "f":
                        formulas[header] = value
                        value = cached_cell.value
                        if value is None:
                            plan.uncached_formulas += 1
                    if cell.data_type == "e" or cached_cell.data_type == "e":
                        raise ValueError(f"{label}: Excel error in {header}")
                    values[header] = _json_value(value, label)

                sn = values["SN"]
                if isinstance(sn, bool) or sn is None or str(sn).strip().casefold() in {"", "n/a"}:
                    raise ValueError(f"{label}: SN is required")
                if isinstance(sn, float):
                    if not sn.is_integer():
                        raise ValueError(f"{label}: numeric SN must be a whole number")
                    sn = int(sn)
                plot_number = str(sn).strip()
                if len(plot_number) > 64:
                    raise ValueError(f"{label}: SN must be 64 characters or fewer")
                location = LOCATIONS.get(str(values["Location"]).strip().casefold())
                harvest_number = HARVEST_NUMBERS.get(str(values["Harvest"]).strip().casefold())
                if location is None or harvest_number is None:
                    raise ValueError(f"{label}: unrecognized Location or Harvest")
                raw_date = values["Date of harvest"]
                try:
                    harvest_date = date.fromisoformat(raw_date)
                except (TypeError, ValueError) as exc:
                    raise ValueError(f"{label}: Date of harvest must be a recorded date") from exc
                if harvest_date.year != 2026:
                    raise ValueError(f"{label}: this importer expects 2026 harvest dates")

                dynamic_data = {COLUMNS[header]: value for header, value in values.items() if header in COLUMNS}
                dynamic_data["_source"] = {
                    "dataset": DATASET,
                    "file": path.name,
                    "sheet": sheet.title,
                    "row": row_number,
                    "plot_number_from": "SN (temporary, not a physical plot identifier)",
                    "column_map": {COLUMNS[header]: header for header in headers if header in COLUMNS},
                    "formulas": formulas,
                    **index.get(sheet.title, {}),
                }
                plan.rows.append(WorkbookRow(location, harvest_date, harvest_number, plot_number, dynamic_data))
        if not plan.rows:
            raise ValueError("Workbook contains no dated harvest records")
        return plan
    finally:
        workbook.close()
        cached.close()


def _json_value(value, label: str):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError(f"{label}: non-finite numeric value")
    if value is not None and not isinstance(value, (str, int, float, bool)):
        raise ValueError(f"{label}: unsupported cell value")
    return value


def import_workbook(db: Session, plan: WorkbookImport) -> dict[str, int]:
    # Keep two runs of this importer from creating the same records concurrently.
    if db.get_bind().dialect.name == "postgresql":
        db.execute(text("SELECT pg_advisory_xact_lock(:key)"), {"key": 20260024})

    existing = {}
    for record, event, location in db.execute(
        select(HarvestRecord, HarvestEvent, Field)
        .join(HarvestEvent, HarvestRecord.harvest_event_id == HarvestEvent.id)
        .join(Field, HarvestEvent.field_id == Field.id)
    ):
        source = (record.dynamic_data or {}).get("_source", {})
        if not isinstance(source, dict):
            continue
        if source.get("dataset") != DATASET:
            continue
        key = source.get("sheet"), source.get("row")
        if key in existing:
            raise ValueError(f"Duplicate imported source row in database: {key}")
        existing[key] = record, event, location

    new_rows = []
    source_keys = set()
    for row in plan.rows:
        if row.source_key in source_keys:
            raise ValueError(f"Duplicate source row in workbook: {row.source_key}")
        source_keys.add(row.source_key)
        previous = existing.get(row.source_key)
        if previous is None:
            new_rows.append(row)
            continue
        record, event, location = previous
        same_event = (
            location.name.casefold() == row.location.casefold()
            and event.harvest_date == row.harvest_date
            and event.harvest_number == row.harvest_number
        )
        old_data = dict(record.dynamic_data)
        new_data = dict(row.dynamic_data)
        # Renaming the file does not change the identity of its source rows.
        old_data["_source"] = {key: value for key, value in old_data["_source"].items() if key != "file"}
        new_data["_source"] = {key: value for key, value in new_data["_source"].items() if key != "file"}
        if not same_event or record.plot_number != row.plot_number or old_data != new_data:
            raise ValueError(f"Previously imported row changed: {row.source_key}. Review it before importing again.")

    result = {"fields_created": 0, "events_created": 0, "records_created": 0, "records_skipped": len(plan.rows) - len(new_rows)}
    fields = {}
    for location in db.scalars(select(Field)):
        key = location.name.casefold()
        if key in fields:
            raise ValueError(f"Ambiguous field name: {location.name}")
        fields[key] = location
    events = {}
    for event in db.scalars(select(HarvestEvent)):
        key = event.field_id, event.harvest_date, event.harvest_number
        events.setdefault(key, []).append(event)

    for row in new_rows:
        location = fields.get(row.location.casefold())
        if location is None:
            location = Field(name=row.location)
            db.add(location)
            db.flush()
            fields[row.location.casefold()] = location
            result["fields_created"] += 1
        key = location.id, row.harvest_date, row.harvest_number
        matching_events = events.get(key, [])
        if len(matching_events) > 1:
            raise ValueError(f"Ambiguous harvest event for {row.event_key}")
        if matching_events:
            event = matching_events[0]
        else:
            event = HarvestEvent(field_id=location.id, harvest_date=row.harvest_date, harvest_number=row.harvest_number)
            db.add(event)
            db.flush()
            events[key] = [event]
            result["events_created"] += 1
        db.add(HarvestRecord(harvest_event_id=event.id, plot_number=row.plot_number, dynamic_data=row.dynamic_data))
        result["records_created"] += 1
    db.flush()
    return result
