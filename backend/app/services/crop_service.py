import csv
import io
import json
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import ValidationError
from sqlalchemy.orm import Session

from app.models.field import Field
from app.models.harvest_event import HarvestEvent
from app.models.harvest_record import HarvestRecord
from app.schemas.harvest_record import HarvestRecordCreate, HarvestRecordQueryRead, HarvestRecordRead, HarvestRecordUpdate
from app.schemas.harvest_record_filter import HarvestRecordFilter


class CropService:
    def create(
        self, db: Session, payload: HarvestRecordCreate, harvest_event_id: int
    ) -> HarvestRecordRead | None:
        if db.get(HarvestEvent, harvest_event_id) is None:
            return None

        record = HarvestRecord(
            harvest_event_id=harvest_event_id,
            plot_number=payload.plot_number,
            dynamic_data=payload.dynamic_data,
        )
        db.add(record)
        db.commit()
        db.refresh(record)
        return _to_read(record)

    def update(
        self, db: Session, record_id: int, payload: HarvestRecordUpdate
    ) -> HarvestRecordRead | None:
        record = db.get(HarvestRecord, record_id)
        if record is None:
            return None

        updates = payload.model_dump(exclude_unset=True)
        if "plot_number" in updates:
            record.plot_number = updates["plot_number"]
        if "dynamic_data" in updates:
            record.dynamic_data = updates["dynamic_data"]

        db.commit()
        db.refresh(record)
        return _to_read(record)

    def delete(self, db: Session, record_id: int) -> bool:
        record = db.get(HarvestRecord, record_id)
        if record is None:
            return False

        db.delete(record)
        db.commit()
        return True

    def import_csv(self, db: Session, csv_text: str, harvest_event_id: int) -> list[HarvestRecordRead]:
        reader = csv.DictReader(io.StringIO(csv_text))
        if not reader.fieldnames:
            raise ValueError("CSV file must include a header row")

        required_columns = {"plot_number"}
        missing_columns = required_columns - set(reader.fieldnames)
        if missing_columns:
            missing_column_names = ", ".join(sorted(missing_columns))
            raise ValueError(f"Missing required column: {missing_column_names}")

        return self._import_rows(db, enumerate(reader, start=2), harvest_event_id)

    def import_xlsx(self, db: Session, raw: bytes, harvest_event_id: int) -> list[HarvestRecordRead]:
        workbook = None
        try:
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
            if not workbook.worksheets:
                raise ValueError("XLSX file must contain a worksheet")
            rows = workbook.worksheets[0].iter_rows()
            header_cells = next(rows, ())
            headers = [str(cell.value).strip() if cell.value is not None else "" for cell in header_cells]
            if not any(headers):
                raise ValueError("The first worksheet must include a header row")
            if "plot_number" not in headers:
                raise ValueError("Missing required column: plot_number")
            named_headers = [header for header in headers if header]
            if len(named_headers) != len(set(named_headers)):
                raise ValueError("XLSX column names must not be repeated")

            def worksheet_rows():
                for row_number, cells in enumerate(rows, start=2):
                    if all(cell.value is None or cell.value == "" for cell in cells):
                        continue
                    if any(cell.data_type == "f" for cell in cells):
                        raise ValueError(f"Row {row_number}: replace formulas with values before uploading")
                    if any(cell.data_type == "e" for cell in cells):
                        raise ValueError(f"Row {row_number}: fix Excel cell errors before uploading")
                    yield row_number, {header: cell.value for header, cell in zip(headers, cells) if header}

            return self._import_rows(db, worksheet_rows(), harvest_event_id)
        except (BadZipFile, InvalidFileException, ParseError, KeyError, EOFError, OSError) as exc:
            raise ValueError("Could not read XLSX file. Upload a valid, unencrypted .xlsx workbook.") from exc
        finally:
            if workbook is not None:
                workbook.close()

    def _import_rows(self, db: Session, rows, harvest_event_id: int) -> list[HarvestRecordRead]:
        records: list[HarvestRecord] = []
        for line_number, row in rows:
            try:
                data = _clean_row(row)
                if data.get("plot_number") is not None:
                    data["plot_number"] = str(data["plot_number"])
                dynamic_data = data.get("dynamic_data")
                if dynamic_data is None:
                    data.pop("dynamic_data", None)
                else:
                    data["dynamic_data"] = json.loads(dynamic_data)
                payload = HarvestRecordCreate(**data)
            except (json.JSONDecodeError, TypeError) as exc:
                raise ValueError(f"Row {line_number}: dynamic_data must be valid JSON") from exc
            except ValidationError as exc:
                raise ValueError(f"Row {line_number}: {exc}") from exc

            records.append(
                HarvestRecord(
                    harvest_event_id=harvest_event_id,
                    plot_number=payload.plot_number,
                    dynamic_data=payload.dynamic_data,
                )
            )

        db.add_all(records)
        db.commit()
        return [_to_read(record) for record in records]

    def list_records(
        self, db: Session, filters: HarvestRecordFilter
    ) -> tuple[list[HarvestRecordQueryRead], int]:
        query = (
            db.query(HarvestRecord, HarvestEvent.harvest_date, Field.name.label("field_name"))
            .select_from(HarvestRecord)
            .join(HarvestEvent, HarvestRecord.harvest_event_id == HarvestEvent.id)
            .join(Field, HarvestEvent.field_id == Field.id)
        )

        if filters.harvest_event_id is not None:
            query = query.filter(HarvestRecord.harvest_event_id == filters.harvest_event_id)
        if filters.field_id is not None:
            query = query.filter(HarvestEvent.field_id == filters.field_id)
        if filters.plot_number is not None:
            query = query.filter(HarvestRecord.plot_number.ilike(f"%{filters.plot_number}%"))
        if filters.harvest_date_from is not None:
            query = query.filter(HarvestEvent.harvest_date >= filters.harvest_date_from)
        if filters.harvest_date_to is not None:
            query = query.filter(HarvestEvent.harvest_date <= filters.harvest_date_to)

        total_count = query.count()
        query = query.order_by(HarvestRecord.id)
        offset = (filters.page - 1) * filters.page_size
        records = query.offset(offset).limit(filters.page_size).all()

        return [
            HarvestRecordQueryRead(
                **HarvestRecordRead.model_validate(record).model_dump(),
                harvest_date=harvest_date,
                field_name=field_name,
            )
            for record, harvest_date, field_name in records
        ], total_count


def _to_read(record: HarvestRecord) -> HarvestRecordRead:
    return HarvestRecordRead(
        id=record.id,
        harvest_event_id=record.harvest_event_id,
        plot_number=record.plot_number,
        dynamic_data=record.dynamic_data or {},
        created_at=record.created_at,
    )


def _clean_row(row: dict) -> dict:
    cleaned = {}
    for key, value in row.items():
        if key is None:
            continue
        if isinstance(value, str):
            value = value.strip()
        cleaned[key] = None if value == "" else value
    return cleaned
